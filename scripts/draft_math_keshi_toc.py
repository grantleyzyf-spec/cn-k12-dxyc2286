# -*- coding: utf-8 -*-
"""小学 8 册《每节内容对应页码》草稿 + 正文页序索引（v3，最终形态）。

为什么不再强求自动抽课时：实测五上（有教师人工抄录的基准 30 条课时）——字号法最多命中 4 条，
且五上/一上 PDF 里课时标题并不以「大字号文本行」存在（14pt 的短行是「阅读理解/分析解答/回顾反思」这类栏目）。
⇒ 本脚本产出两件**可靠的半成品**，把人工抄录变成「照着索引填表」：

  ① `…每节内容对应页码（草稿·待校）.txt`
     单元骨架（按图谱 chapter，即教材单元顺序）+ 单元页码范围占位 + 有 docx 时的课时候选（带页码估算）
  ② `…正文页序索引.txt`
     每页一行：`教材页 <pp>　(PDF p<i>)　首行文本…`，人工对着它抄课时标题最快

用法：python3 draft_keshi_toc_v3.py [--book 二年级上册] [--write]
"""
import argparse
import glob
import json
import os
import re
from collections import Counter

import fitz

XI = "/Users/a627692/Desktop/中小学教材/中小学教材/数学/小学/人教版"
REPO = os.path.expanduser("~/Desktop/cn-k12-dxyc2286")
BOOKS = ["一年级上册", "一年级下册", "二年级上册", "二年级下册",
         "三年级上册", "三年级下册", "四年级上册", "六年级上册"]
CN = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]
GD = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6}
NOISE = re.compile(r"仅供个人|未经授权|不得另做他用|电子课本|公众号|微信|淘宝|义务\s*教\s*育\s*教\s*科\s*书")
NAME_RE = [
    re.compile(r"第\s*([一二三四五六七八九十\d]+)\s*课时[，、：:．.\s]*([^\s，。；、（）()「」“”]{2,14})"),
]
BAD_NAME = re.compile(r"^(在|以|为|先|边|通过|重在|学习|理解|掌握|认识|感受|经历|体会|能|会|使|让)")


def norm(s):
    return re.sub(r"[\s　（）()【】\[\]、，,．.·・「」“”\"'’!！?？:：;；—\-]", "", s or "")


def find_pdf(book):
    c = [p for p in glob.glob(f"{XI}/*.pdf") if book[:3] in p and book[-2:] in p]
    if not c:
        c = [p for p in glob.glob(f"{XI}/*.pdf") if book[:3] in p]
    return c[0] if c else None


def page_lines(page):
    out = []
    for b in page.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            t = "".join(s["text"] for s in l["spans"]).strip()
            if t and not NOISE.search(t):
                out.append((round(max(s["size"] for s in l["spans"]), 1), t))
    return out


def cand_from_docx(book):
    """有《大单元整体教学设计》的册：提「第N课时 X」候选（单元名来自文件名）。"""
    out = {}
    try:
        from docx import Document
    except ImportError:
        return out
    docs = (glob.glob(f"{XI}/{book}/教学设计/*大单元整体教学设计*.docx") +
            glob.glob(f"{XI}/{book}/教学设计/*.docx") +
            glob.glob(f"{XI}/{book}/单元作业设计/*.docx") +
            glob.glob(f"{XI}/{book}/*单元作业设计*.docx"))
    for p in docs:
        m = re.search(r"第\s*([一二三四五六七八九十\d]+)\s*单元\s*([^\s\[（]+)", os.path.basename(p))
        if not m:
            continue
        uname = m.group(2).strip()
        try:
            d = Document(p)
        except Exception:
            continue
        txt = "\n".join(x.text for x in d.paragraphs)
        for tb in d.tables:
            for r in tb.rows:
                txt += "\n" + " ".join(c.text for c in r.cells)
        names = []
        for rx in NAME_RE:
            for mm in rx.finditer(txt):
                nm = re.sub(r"\s+", "", mm.group(2)).strip("，。、；：")
                if 2 <= len(norm(nm)) <= 14 and not BAD_NAME.match(nm) and not re.search(r"[0-9０-９=＝+＋]", nm):
                    names.append(nm)
        seen = set()
        for nm in names:
            if norm(nm) in seen:
                continue
            seen.add(norm(nm))
            out.setdefault(uname, []).append(nm)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--book", action="append", default=None)
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    topics = json.load(open(f"{REPO}/data/math/topics.json", encoding="utf-8"))

    for book in a.book or BOOKS:
        pdf = find_pdf(book)
        if not pdf:
            print(f"✗ {book} 缺 PDF")
            continue
        d = fitz.open(pdf)
        pages = [page_lines(d[i]) for i in range(d.page_count)]
        offs = []
        for i, ls in enumerate(pages):
            for _sz, t in reversed(ls[-3:]):
                if re.fullmatch(r"\d{1,3}", t):
                    offs.append(int(t) - (i + 1))
                    break
        offset = Counter(offs).most_common(1)[0][0] if offs else 0

        g, sem = GD[book[0]], ("上" if book.endswith("上册") else "下")
        units, seen = [], set()
        for t in topics:
            if (t["grade"], t["semester"]) != (g, sem):
                continue
            u = re.sub(r"^第[一二三四五六七八九十]+单元\s*", "", t["chapter"].split(" ", 1)[1])
            if u not in seen:
                seen.add(u)
                units.append(u)

        # ① 草稿
        cand = cand_from_docx(book)
        body, filled, review = [], 0, []
        for k, uname in enumerate(units):
            key = next((u for u in cand if norm(u) == norm(uname) or norm(u) in norm(uname)
                        or norm(uname) in norm(u)), None)
            names = cand.get(key, []) if key else []
            body.append(f"{CN[k] if k < 10 else k+1}　{uname}（第?—?页）")
            if not names:
                body.append("　（课时待抄录：见同目录《正文页序索引》）")
                review.append(f"第{k+1}单元 {uname}：需人工抄录")
            else:
                for nm in names:
                    key_n = norm(nm)
                    pg, head = None, False
                    for i, ls in enumerate(pages):
                        joined = norm("".join(t for _sz, t in ls))
                        if key_n and key_n in joined:
                            head = key_n in norm("".join(t for _sz, t in ls[:5]))
                            pg = i
                            if head:
                                break
                    if pg is None:
                        body.append(f"　{nm} ?（PDF 未检出）")
                        review.append(f"{uname}：`{nm}` 未检出")
                    else:
                        body.append(f"　{nm} {pg + 1 + offset}{'' if head else '?'}")
                        filled += 1
            body.append("")

        draft = [f"数学（人教版）{book}　每节内容对应页码（草稿·待校）",
                 f"（教材页码＝PDF 页 {offset:+d} 修正；单元 {len(units)} 个；本册课时行 {filled} 条已带页码线索）",
                 "",
                 "【怎么用】",
                 "　1) 打开同目录《数学人教版%s正文页序索引.txt》，逐页扫一遍，把每个课时标题抄进本文件对应单元下；" % book,
                 "　2) 每行格式：`　课时名 教材页码`（缩进两个全角空格）；单元标题行写成 `一　单元名（第a—b页）`；",
                 "　3) 带 `?` 的页码表示该名字只在正文中段检出、不是课时起始页，必须核；",
                 "　4) 校完后把文件名里的「（草稿·待校）」去掉，即可作为 audit_math_yijiao_sections.py 的课时来源。",
                 ""]
        if review:
            draft += ["【待校清单】"] + [f"　· {r}" for r in review] + [""]
        draft += ["━━━━━━━━━━━━━━━━━━━━━━"] + body

        # ② 页序索引
        idx = [f"数学（人教版）{book}　正文页序索引（用于人工抄录课时标题）",
               f"（每页一行：教材页 / PDF 页 / 该页最上方两行文本；页脚偏移 {offset:+d}）", ""]
        for i, ls in enumerate(pages):
            head = " ／ ".join(t for _sz, t in ls[:2]) or "（无文本）"
            idx.append(f"教材页 {i + 1 + offset:>3}　(PDF p{i + 1:>3})　{head[:60]}")

        outdir = os.path.join(XI, book, "教学设计")
        p1 = os.path.join(outdir, f"数学人教版{book}每节内容对应页码（草稿·待校）.txt")
        p2 = os.path.join(outdir, f"数学人教版{book}正文页序索引.txt")
        print(f"{book}: 单元 {len(units)} · 课时行 {filled} · 待校 {len(review)} · 索引 {len(pages)} 页 → 教学设计/")
        if a.write:
            os.makedirs(outdir, exist_ok=True)
            open(p1, "w", encoding="utf-8").write("\n".join(draft).rstrip() + "\n")
            open(p2, "w", encoding="utf-8").write("\n".join(idx).rstrip() + "\n")


if __name__ == "__main__":
    main()
